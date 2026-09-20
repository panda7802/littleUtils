#!/usr/bin/env python3
"""Collect public posts from a Douyin account with a real browser session.

This intentionally does not reverse-engineer Douyin's request signatures.  A
normal Chromium page makes the signed requests; the script observes the public
post-list responses while scrolling the profile page.
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from playwright.sync_api import Page, Response, TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


POST_API_MARKERS = ("/aweme/v1/web/aweme/post/",)


def clean_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    return value[:80] or "douyin_account"


def first_url(value: Any) -> str:
    if not isinstance(value, dict):
        return ""
    urls = value.get("url_list")
    return str(urls[0]) if isinstance(urls, list) and urls else ""


def post_to_row(post: dict[str, Any]) -> dict[str, Any]:
    aweme_id = str(post.get("aweme_id") or post.get("group_id") or "")
    content = str(post.get("desc") or "").strip()
    title = str(post.get("preview_title") or post.get("item_title") or "").strip()
    if not title:
        title = next((line.strip() for line in content.splitlines() if line.strip()), "")
    title = title[:100]

    video = post.get("video") if isinstance(post.get("video"), dict) else {}
    play_url = first_url(video.get("play_addr")) or first_url(video.get("download_addr"))
    cover_url = (
        first_url(video.get("cover"))
        or first_url(video.get("origin_cover"))
        or first_url(video.get("dynamic_cover"))
    )

    image_urls: list[str] = []
    for image in post.get("images") or []:
        if not isinstance(image, dict):
            continue
        urls = image.get("url_list")
        url = str(urls[0]) if isinstance(urls, list) and urls else ""
        if url:
            image_urls.append(url)

    created = post.get("create_time")
    created_at: datetime | None = None
    if isinstance(created, (int, float)):
        # Excel does not support timezone-aware datetime values.
        created_at = datetime.fromtimestamp(created)

    author = post.get("author") if isinstance(post.get("author"), dict) else {}
    return {
        "aweme_id": aweme_id,
        "author": author.get("nickname") or "",
        "douyin_id": author.get("unique_id") or author.get("short_id") or "",
        "created_at": created_at,
        "title": title,
        "content": content,
        "content_type": "images" if image_urls else "video",
        "video_page_url": (
            f"https://www.douyin.com/{'note' if image_urls else 'video'}/{aweme_id}"
            if aweme_id else ""
        ),
        # CDN URLs are temporary and may require the same browser cookies/referrer.
        "play_url": play_url,
        "cover_url": cover_url,
        "image_urls": image_urls,
    }


def find_profile_url(page: Page, account: str, login_wait: int) -> str:
    search_url = f"https://www.douyin.com/search/{quote(account)}?type=user"
    print(f"[1/3] 搜索抖音号 {account} ...")
    page.goto(search_url, wait_until="domcontentloaded", timeout=60_000)
    deadline = time.monotonic() + login_wait
    warned = False
    while True:
        page.wait_for_timeout(2_000)
        candidates = page.locator('a[href*="/user/"]')
        for index in range(min(candidates.count(), 30)):
            anchor = candidates.nth(index)
            href = anchor.get_attribute("href") or ""
            try:
                context = anchor.evaluate(
                    "el => (el.closest('[data-e2e]') || el.parentElement?.parentElement || el).innerText"
                )
            except Exception:
                context = anchor.inner_text()
            if account in (context or ""):
                if href.startswith("//"):
                    return "https:" + href
                if href.startswith("/"):
                    return "https://www.douyin.com" + href
                return href
        if time.monotonic() >= deadline:
            break
        if not warned:
            print(f"      若页面要求登录或验证码，请在浏览器中完成（最多等待 {login_wait} 秒）")
            warned = True

    raise RuntimeError(
        "搜索结果中没有找到完全匹配的抖音号。可能需要先登录/完成验证码，"
        "也可以从浏览器复制该用户主页 URL，再用 --profile-url 传入。"
    )


def collect_posts(page: Page, profile_url: str, max_posts: int, max_idle: int) -> list[dict[str, Any]]:
    posts: dict[str, dict[str, Any]] = {}
    last_change = time.monotonic()

    def on_response(response: Response) -> None:
        nonlocal last_change
        if not any(marker in response.url for marker in POST_API_MARKERS):
            return
        try:
            payload = response.json()
        except Exception:
            return
        items = payload.get("aweme_list") or payload.get("data", {}).get("aweme_list") or []
        before = len(posts)
        for item in items:
            if not isinstance(item, dict):
                continue
            item_id = str(item.get("aweme_id") or item.get("group_id") or "")
            if item_id:
                posts[item_id] = item
        if len(posts) > before:
            last_change = time.monotonic()
            print(f"      已获取 {len(posts)} 条作品")

    page.on("response", on_response)
    print(f"[2/3] 打开主页并滚动采集：{profile_url}")
    page.goto(profile_url, wait_until="domcontentloaded", timeout=60_000)
    page.wait_for_timeout(5_000)

    unchanged_height_rounds = 0
    previous_height = 0
    while len(posts) < max_posts:
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        page.wait_for_timeout(1_500)
        height = page.evaluate("document.body.scrollHeight")
        unchanged_height_rounds = unchanged_height_rounds + 1 if height == previous_height else 0
        previous_height = height

        idle_seconds = time.monotonic() - last_change
        if idle_seconds >= max_idle and unchanged_height_rounds >= 3:
            break

    page.remove_listener("response", on_response)
    rows = [post_to_row(post) for post in posts.values()]
    rows.sort(key=lambda row: row["created_at"] or datetime.min, reverse=True)
    return rows[:max_posts]


def write_results(rows: list[dict[str, Any]], output_dir: Path, account: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    author_name = next((str(row.get("author") or "").strip() for row in rows if row.get("author")), "")
    xlsx_path = output_dir / f"{clean_filename(author_name or account)}.xlsx"

    columns = [
        ("aweme_id", "作品ID"),
        ("author", "作者昵称"),
        ("douyin_id", "抖音号"),
        ("created_at", "发布时间"),
        ("title", "标题"),
        ("content", "内容"),
        ("content_type", "作品类型"),
        ("video_page_url", "作品链接"),
        ("play_url", "播放地址（临时）"),
        ("cover_url", "封面地址"),
        ("image_urls", "图片地址"),
    ]
    type_names = {"video": "视频", "images": "图文"}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "作品列表"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append([label for _, label in columns])

    for row in rows:
        values: list[Any] = []
        for key, _ in columns:
            value = row.get(key, "")
            if key == "content_type":
                value = type_names.get(str(value), value)
            elif key == "image_urls":
                value = "\n".join(str(url) for url in (value or []))
            values.append(value)
        sheet.append(values)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10, color="222222")
    thin_gray = Side(style="thin", color="D9E2F3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin_gray)
    sheet.row_dimensions[1].height = 24

    for row_index in range(2, sheet.max_row + 1):
        for column_index in range(1, len(columns) + 1):
            cell = sheet.cell(row_index, column_index)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=column_index in {5, 6, 11})
        sheet.cell(row_index, 4).number_format = "yyyy-mm-dd hh:mm:ss"
        for column_index in (8, 9, 10):
            cell = sheet.cell(row_index, column_index)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    widths = [22, 18, 16, 20, 32, 60, 12, 42, 42, 42, 42]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    if rows:
        table = Table(displayName="DouyinPosts", ref=f"A1:K{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = "A1:K1"

    workbook.save(xlsx_path)
    return xlsx_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="采集抖音公开账号的作品链接、标题和文案")
    parser.add_argument("account", help="抖音号，例如 837672563")
    parser.add_argument("--profile-url", help="已知用户主页 URL 时建议直接传入，最稳定")
    parser.add_argument("--max-posts", type=int, default=10_000, help="最多采集作品数")
    parser.add_argument("--max-idle", type=int, default=12, help="多少秒无新数据后停止")
    parser.add_argument("--login-wait", type=int, default=120, help="等待手动登录/验证码的秒数")
    parser.add_argument("--output", type=Path, default=Path("output"), help="输出目录")
    parser.add_argument("--headless", action="store_true", help="无头运行；首次使用不建议开启")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.max_posts < 1 or args.max_idle < 3 or args.login_wait < 0:
        print("--max-posts 必须 >= 1，--max-idle 必须 >= 3，--login-wait 必须 >= 0", file=sys.stderr)
        return 2

    session_dir = Path(".douyin-browser").resolve()
    try:
        with sync_playwright() as playwright:
            context = playwright.chromium.launch_persistent_context(
                str(session_dir),
                headless=args.headless,
                channel="chrome",
                viewport={"width": 1440, "height": 900},
                locale="zh-CN",
            )
            page = context.pages[0] if context.pages else context.new_page()
            login_wait = min(args.login_wait, 10) if args.headless else args.login_wait
            profile_url = args.profile_url or find_profile_url(page, args.account, login_wait)
            rows = collect_posts(page, profile_url, args.max_posts, args.max_idle)
            context.close()
    except PlaywrightTimeoutError as exc:
        print(f"页面加载超时：{exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"采集失败：{exc}", file=sys.stderr)
        print("请确认浏览器中已登录且没有未完成的验证码。", file=sys.stderr)
        return 1

    xlsx_path = write_results(rows, args.output, args.account)
    print(f"[3/3] 完成，共 {len(rows)} 条")
    print(f"      Excel: {xlsx_path.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
