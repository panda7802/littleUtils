"""Shared XLSX output helpers for social-media collectors."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


Column = tuple[str, str, int]


def clean_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")
    return value[:80] or "账号作品"


def write_social_xlsx(
    rows: list[dict[str, Any]],
    output_dir: Path,
    display_name: str,
    fallback_name: str,
    columns: list[Column],
    *,
    sheet_name: str,
    table_name: str,
    date_keys: Iterable[str] = (),
    hyperlink_keys: Iterable[str] = (),
    wrap_keys: Iterable[str] = (),
    transforms: dict[str, Callable[[Any], Any]] | None = None,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{clean_filename(display_name or fallback_name)}.xlsx"
    date_keys = set(date_keys)
    hyperlink_keys = set(hyperlink_keys)
    wrap_keys = set(wrap_keys)
    transforms = transforms or {}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append([label for _, label, _ in columns])

    for row in rows:
        values: list[Any] = []
        for key, _, _ in columns:
            value = row.get(key, "")
            if key in transforms:
                value = transforms[key](value)
            values.append(value)
        sheet.append(values)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10, color="222222")
    thin_gray = Side(style="thin", color="D9E2F3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin_gray)
    sheet.row_dimensions[1].height = 24

    for row_index in range(2, sheet.max_row + 1):
        for column_index, (key, _, _) in enumerate(columns, start=1):
            cell = sheet.cell(row_index, column_index)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=key in wrap_keys)
            if key in date_keys and cell.value:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if key in hyperlink_keys and cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"

    for column_index, (_, _, width) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    final_column = get_column_letter(len(columns))
    if rows:
        table = Table(displayName=table_name, ref=f"A1:{final_column}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = f"A1:{final_column}1"

    workbook.save(output_path)
    return output_path
